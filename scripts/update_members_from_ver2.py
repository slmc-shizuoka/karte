import json
from collections import OrderedDict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path('/Users/owner/Downloads/会員データver2.xlsx')
TARGET = ROOT / 'dist' / 'data' / 'members.json'


def text(value):
    return '' if value is None else str(value).strip()


existing = json.loads(TARGET.read_text(encoding='utf-8'))
by_id = {text(member['memberId']): member for member in existing}

workbook = load_workbook(SOURCE, read_only=True, data_only=True)
sheet = workbook.active
members = OrderedDict()

for member_id, postal_code, _prefecture, address, status in sheet.iter_rows(min_row=2, values_only=True):
    member_id = text(member_id)
    if not member_id:
        continue
    if member_id not in by_id:
        raise ValueError(f'地区グループが未登録の会員番号です: {member_id}')
    if member_id not in members:
        previous = by_id[member_id]
        members[member_id] = {
            'memberId': member_id,
            'postalCode': text(postal_code),
            'districtGroup': previous['districtGroup'],
            'area': previous.get('area') or text(address).split('【', 1)[0],
            'status': [],
            'shelf': previous.get('shelf', 'UNKNOWN'),
        }
    member_status = text(status)
    if member_status and member_status not in members[member_id]['status']:
        members[member_id]['status'].append(member_status)

if set(members) != set(by_id):
    missing = sorted(set(by_id) - set(members))
    raise ValueError(f'新ファイルに存在しない会員番号があります: {missing[:10]}')

result = []
for member in members.values():
    member['status'] = '・'.join(member['status']) or '未設定'
    result.append(member)

TARGET.write_text(json.dumps(result, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
print(json.dumps({
    'members': len(result),
    'multi_status_members': sum('・' in member['status'] for member in result),
    'statuses': sorted({status for member in result for status in member['status'].split('・')}),
}, ensure_ascii=False))
